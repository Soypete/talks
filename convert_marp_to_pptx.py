from pathlib import Path
import re
import subprocess

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).parent
SOURCE = Path('/tmp/agentcon-self-hosting.md')
OUT = ROOT / 'AGNTCon_MCPCon_North_America_2026_full.pptx'
BG = Path('/tmp/agntcon_media/ppt/media/image1.jpg')

commit = subprocess.check_output(
    ['git', 'log', '-1', '--format=%H', '--all', '--', 'agentcon-self-hosting/talk.md'],
    cwd=ROOT, text=True,
).strip()
SOURCE.write_text(subprocess.check_output(
    ['git', 'show', f'{commit}:agentcon-self-hosting/talk.md'],
    cwd=ROOT, text=True,
))

raw = SOURCE.read_text()
parts = re.split(r'^---\s*$', raw, flags=re.M)
slides = [p.strip() for p in parts[2:] if p.strip()]

p = Presentation()
p.slide_width = Inches(13.333)
p.slide_height = Inches(7.5)
blank = p.slide_layouts[6]

NAVY = RGBColor(31, 33, 58)
PANEL = RGBColor(45, 47, 78)
WHITE = RGBColor(248, 248, 252)
MUTED = RGBColor(182, 185, 212)
CYAN = RGBColor(74, 196, 239)
PINK = RGBColor(236, 105, 154)
PURPLE = RGBColor(155, 122, 222)
GREEN = RGBColor(94, 218, 170)


def add_text(slide, value, x, y, w, h, size=16, color=WHITE, bold=False, font='Aptos'):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.06)
    tf.margin_top = tf.margin_bottom = Inches(.03)
    p0 = tf.paragraphs[0]
    run = p0.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_background(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    if BG.exists():
        slide.shapes.add_picture(str(BG), 0, Inches(4.9), width=p.slide_width, height=Inches(2.6))
    cover = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, p.slide_width, Inches(6.2))
    cover.fill.solid(); cover.fill.fore_color.rgb = NAVY; cover.fill.transparency = 5
    cover.line.fill.background()


def clean_inline(s):
    s = re.sub(r'!\[[^]]*\]\(([^)]+)\)', r'[visual: \1]', s)
    s = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', s)
    s = s.replace('**', '').replace('__', '').replace('`', '')
    return s


def parse_section(section):
    lines = section.splitlines()
    heading = next((re.sub(r'^#+\s*', '', x).strip() for x in lines if re.match(r'^#', x)), 'Untitled')
    content = []
    in_code = False
    code = []
    for line in lines:
        if line.strip().startswith('<!--'):
            continue
        if line.strip().startswith('```'):
            if in_code:
                content.append(('code', '\n'.join(code)))
                code = []
            in_code = not in_code
            continue
        if in_code:
            code.append(line)
            continue
        if line.strip().startswith('#'):
            continue
        if not line.strip():
            continue
        if line.startswith('|'):
            content.append(('table', clean_inline(line.strip().strip('|'))))
        elif line.lstrip().startswith(('-', '*')):
            content.append(('bullet', clean_inline(line.lstrip()[1:].strip())))
        else:
            content.append(('text', clean_inline(line.strip())))
    return heading, content


for index, section in enumerate(slides, 1):
    slide = p.slides.add_slide(blank)
    add_background(slide)
    heading, content = parse_section(section)
    lead = '<!-- _class: lead -->' in section
    add_text(slide, f'{index:02d}', .65, .28, .5, .3, 11, CYAN, True)
    add_text(slide, heading, .65, .65, 12, 1.0, 28 if lead else 25, WHITE, True)
    y = 1.75 if not lead else 2.05
    for kind, value in content:
        if kind == 'code':
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.8), Inches(y), Inches(11.75), Inches(min(2.2, .48 + .28 * value.count('\n'))))
            box.fill.solid(); box.fill.fore_color.rgb = PANEL; box.line.color.rgb = PURPLE
            add_text(slide, value, 1.0, y + .14, 11.35, min(1.9, .25 + .28 * value.count('\n')), 14, CYAN, False, 'Courier New')
            y += min(2.35, .68 + .28 * value.count('\n'))
        elif kind == 'bullet':
            add_text(slide, '• ' + value, 1.0, y, 11.3, .42, 16, WHITE)
            y += .48
        elif kind == 'table':
            add_text(slide, value, 1.0, y, 11.3, .35, 14, MUTED, False, 'Courier New')
            y += .38
        else:
            color = PINK if value.startswith('[visual:') else WHITE
            size = 19 if lead else 16
            add_text(slide, value, 1.0, y, 11.3, .68, size, color, value.startswith('**'))
            y += .76
    if y > 6.65:
        add_text(slide, 'Content preserved from Marp source; adjust layout in PowerPoint as needed.', .7, 6.85, 11.8, .25, 9, MUTED)

p.core_properties.title = 'Self-Hosting Agents: Small Dense Models Change the Economics'
p.core_properties.subject = 'Full Marp-to-PowerPoint conversion'
p.core_properties.author = 'Pedro / Soypete'
p.save(OUT)
print(f'created {OUT} with {len(p.slides)} slides')
